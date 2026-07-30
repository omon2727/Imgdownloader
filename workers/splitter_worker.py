import os
from qtpy.QtCore import QThread, Signal
import pandas as pd
from openpyxl import load_workbook


class SplitterWorker(QThread):
    progress = Signal(int)
    log = Signal(str)
    finished = Signal()
    result = Signal(int)

    def __init__(self, input_file, output_dir, chunk_size):
        super().__init__()
        self.input_file = input_file
        self.output_dir = output_dir
        self.chunk_size = chunk_size
        self.is_cancelled = False

    def cancel(self):
        self.is_cancelled = True

    def run(self):
        try:
            os.makedirs(self.output_dir, exist_ok=True)
            self.log.emit(f"Открываем файл: {self.input_file}")

            wb = load_workbook(filename=self.input_file, read_only=True)
            ws = wb.active
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

            original_name = os.path.splitext(os.path.basename(self.input_file))[0]
            total_rows = ws.max_row - 1
            total_parts = max(1, (total_rows + self.chunk_size - 1) // self.chunk_size)

            self.log.emit(f"Всего строк: {total_rows:,} | Будет создано частей: {total_parts}")

            row_start = 2
            file_index = 1
            files_created = 0

            while row_start <= ws.max_row:
                if self.is_cancelled:
                    self.log.emit("Разбивка отменена пользователем.")
                    break

                row_end = min(row_start + self.chunk_size - 1, ws.max_row)

                data = []
                for row in ws.iter_rows(min_row=row_start, max_row=row_end, values_only=True):
                    data.append(row)

                df = pd.DataFrame(data, columns=headers)
                output_file = os.path.join(self.output_dir, f'{original_name}_{file_index:03d}.xlsx')
                df.to_excel(output_file, index=False)

                self.log.emit(f'✓ Создан: {os.path.basename(output_file)} — {len(df):,} строк')

                progress_value = int((file_index / total_parts) * 100)
                self.progress.emit(min(progress_value, 100))

                row_start = row_end + 1
                file_index += 1
                files_created += 1

            wb.close()
            self.result.emit(files_created)
            self.log.emit("\n✅ Разбивка успешно завершена!")

        except Exception as e:
            self.log.emit(f"Ошибка: {str(e)}")
        finally:
            self.finished.emit()