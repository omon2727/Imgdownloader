import sys
import os
from qtpy.QtWidgets import (
    QWidget, QPushButton, QVBoxLayout, QHBoxLayout,
    QFileDialog, QMessageBox, QTableView, QLabel, QMenu, QTabWidget,
    QRadioButton, QButtonGroup, QLineEdit, QTextEdit, QProgressBar,
    QSpinBox, QFrame
)
from qtpy.QtCore import QAbstractTableModel, Qt, QThread, Signal
from qtpy.QtGui import QColor, QPixmap, QFont, QIcon
from table_reader import UniversalTableReader
from files_compressor import Compressor
import asyncio
from files_downloader import FileDownloader
import pandas as pd
from openpyxl import load_workbook
from pathlib import Path
from datetime import datetime

# ====================== ВСПОМОГАТЕЛЬНАЯ ФУНКЦИЯ ======================
def get_resource_path(relative_path):
    """Возвращает путь к файлу — работает и в .exe, и при обычном запуске"""
    if hasattr(sys, '_MEIPASS'):
        # Запущено из PyInstaller
        return os.path.join(sys._MEIPASS, relative_path)
    return relative_path
# =====================================================================


# ============================================================================
# MODEL
# ============================================================================
class PandasModel(QAbstractTableModel):
    def __init__(self, df):
        super().__init__()
        self._df = df
        self.url_col = None
        self.name_col = None
        self.brand_col = None

    def rowCount(self, parent=None):
        return self._df.shape[0]

    def columnCount(self, parent=None):
        return self._df.shape[1]

    def data(self, index, role=Qt.DisplayRole):
        if role == Qt.DisplayRole:
            return str(self._df.iloc[index.row(), index.column()])
        if role == Qt.BackgroundRole:
            if index.column() == self.url_col:
                return QColor("#cce5ff")
            if index.column() == self.name_col:
                return QColor("#d4edda")
            if index.column() == self.brand_col:
                return QColor("#fff3cd")
        return None

    def headerData(self, section, orientation, role):
        if role == Qt.DisplayRole and orientation == Qt.Horizontal:
            name = str(self._df.columns[section])
            if section == self.url_col:
                return f"{name} [URL]"
            if section == self.name_col:
                return f"{name} [NAME]"
            if section == self.brand_col:
                return f"{name} [BRAND]"
            return name
        return None

    # ============================================================================
# CONVERTER WORKER (НОВЫЙ)
# ============================================================================
class ConverterWorker(QThread):
    progress = Signal(int)
    log = Signal(str)
    finished = Signal()
    result = Signal(int)

    def __init__(self, source_folder, format_choice):
        super().__init__()
        self.source_folder = source_folder
        self.format_choice = format_choice  # "1" или "2"
        self.is_cancelled = False

    def cancel(self):
        self.is_cancelled = True

    def run(self):
        try:
            source_path = Path(self.source_folder)
            if self.format_choice == "1":
                file_ext = ".txt"
                format_name = "Текст Юникод"
                sep = '\t'
                encoding = 'utf-16'
                lineterminator = '\r\n'
            else:
                file_ext = ".csv"
                format_name = "CSV (разделитель ;)"
                sep = ';'
                encoding = 'utf-8-sig'
                lineterminator = '\r\n'

            timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            output_folder = source_path / f"{format_name.replace(' ', '_')}_{timestamp}"
            output_folder.mkdir(parents=True, exist_ok=True)

            self.log.emit(f"Создана папка: {output_folder}\n")

            excel_files = list(source_path.glob("*.xlsx")) + list(source_path.glob("*.XLSX"))
            excel_files = sorted(set(excel_files))  # уникальные

            self.log.emit(f"Найдено Excel-файлов: {len(excel_files)}\n")

            for i, file_path in enumerate(excel_files, 1):
                if self.is_cancelled:
                    self.log.emit("Конвертация отменена.")
                    break

                try:
                    self.log.emit(f"Обрабатываю: {file_path.name}")
                    df = pd.read_excel(file_path, sheet_name=0)

                    output_filename = file_path.stem + file_ext
                    output_path = output_folder / output_filename

                    df.to_csv(output_path, sep=sep, index=False, encoding=encoding, lineterminator=lineterminator)

                    self.log.emit(f" ✓ Готово: {output_filename} ({len(df)} строк)")
                    self.progress.emit(int((i / len(excel_files)) * 100))

                except Exception as e:
                    self.log.emit(f" ✗ Ошибка {file_path.name}: {e}")

            self.result.emit(len(excel_files))
            self.log.emit("\n✅ Конвертация завершена!")
            self.log.emit(f"Результаты в: {output_folder}")

        except Exception as e:
            self.log.emit(f"Критическая ошибка: {e}")
        finally:
            self.finished.emit()

# ============================================================================
# DOWNLOAD WORKER
# ============================================================================
class DownloadWorker(QThread):
    progress = Signal(int)
    log = Signal(str)
    finished = Signal()
    log_item = Signal(str, bool)
    result = Signal(int, int)
    cancelled = Signal()

    def __init__(self, data, download_dir):
        super().__init__()
        self.data = data
        self.download_dir = download_dir
        self.downloader = None
        self.is_running = False
        self.was_cancelled = False

    def run(self):
        self.is_running = True
        try:
            asyncio.run(self.async_run())
        finally:
            self.is_running = False
            self.finished.emit()

    async def async_run(self):
        self.downloader = FileDownloader()
        self.downloader.set_download_dir(self.download_dir)
        total = len(self.data)
        self.downloader.log_callback = lambda url, ok: self.log_item.emit(url, ok)

        async def monitor():
            while (self.downloader.process_items_count < total and not self.downloader.is_cancelled):
                self.progress.emit(self.downloader.process_items_count)
                await asyncio.sleep(0.1)

        monitor_task = asyncio.create_task(monitor())

        try:
            await self.downloader.download_files(self.data)
        except asyncio.CancelledError:
            pass
        finally:
            if not monitor_task.done():
                monitor_task.cancel()
                try:
                    await monitor_task
                except asyncio.CancelledError:
                    pass
            self.progress.emit(self.downloader.process_items_count)
            total = self.downloader.total_items_count
            errors_count = len(self.downloader.errors)

            if self.downloader.is_cancelled:
                self.was_cancelled = True
                self.cancelled.emit()
            else:
                self.result.emit(total, errors_count)

            if self.downloader.errors:
                for err in self.downloader.errors:
                    self.log.emit(err)

    def cancel_download(self):
        if self.downloader is not None:
            self.downloader.cancel()


# ============================================================================
# COMPRESSOR WORKER
# ============================================================================
class CompressorWorker(QThread):
    progress = Signal(int)
    log = Signal(str)
    finished = Signal()
    result = Signal(int, int)
    cancelled = Signal()

    def __init__(self, compressor: Compressor):
        super().__init__()
        self.compressor = compressor
        self.is_running = False
        self.was_cancelled = False

    def run(self):
        self.is_running = True
        try:
            self.compressor.log_callback = self._on_file_packed
            self.compressor.run()
            if self.compressor._cancelled:
                self.was_cancelled = True
                self.cancelled.emit()
            else:
                self.log.emit("Архивирование завершено успешно!")
                self.result.emit(self.compressor.all_files_count, 0)
        except Exception as e:
            self.log.emit(f"Ошибка при архивировании: {str(e)}")
            self.result.emit(0, 1)
        finally:
            self.is_running = False
            self.finished.emit()

    def _on_file_packed(self, file_name: str):
        self.progress.emit(self.compressor.files_count)
        self.log.emit(f"Добавлено в архив: {file_name}")

    def cancel_compression(self):
        if self.compressor:
            self.compressor.cancel()


# ============================================================================
# SPLITTER WORKER (НОВЫЙ)
# ============================================================================
class SplitterWorker(QThread):
    progress = Signal(int)
    log = Signal(str)
    finished = Signal()
    result = Signal(int)

    def __init__(self, input_file, output_dir, chunk_size):
        super().__init__()
        self.input_file = input_file
        self.output_dir = output_dir
        self.chunk_size = chunk_size
        self.is_cancelled = False

    def cancel(self):
        self.is_cancelled = True

    def run(self):
        try:
            os.makedirs(self.output_dir, exist_ok=True)
            self.log.emit(f"Открываем файл: {self.input_file}")

            wb = load_workbook(filename=self.input_file, read_only=True)
            ws = wb.active
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

            original_name = os.path.splitext(os.path.basename(self.input_file))[0]
            total_rows = ws.max_row - 1  # без заголовка
            total_parts = (total_rows + self.chunk_size - 1) // self.chunk_size

            self.log.emit(f"Всего строк: {total_rows:,} | Будет создано частей: {total_parts}")

            row_start = 2
            file_index = 1
            files_created = 0

            while row_start <= ws.max_row:
                if self.is_cancelled:
                    self.log.emit("Разбивка отменена пользователем.")
                    break

                row_end = min(row_start + self.chunk_size - 1, ws.max_row)

                data = []
                for row in ws.iter_rows(min_row=row_start, max_row=row_end, values_only=True):
                    data.append(row)

                df = pd.DataFrame(data, columns=headers)

                output_file = os.path.join(self.output_dir, f'{original_name}_{file_index:03d}.xlsx')
                df.to_excel(output_file, index=False)

                self.log.emit(f'✓ Создан: {os.path.basename(output_file)} — {len(df):,} строк')

                # Правильный процент
                progress_value = int((file_index / total_parts) * 100)
                self.progress.emit(progress_value)

                row_start = row_end + 1
                file_index += 1
                files_created += 1

            wb.close()
            self.result.emit(files_created)
            self.log.emit("\n✅ Разбивка успешно завершена!")

        except Exception as e:
            self.log.emit(f"Ошибка: {str(e)}")
        finally:
            self.finished.emit()


# ============================================================================
# MAIN WINDOW
# ============================================================================
class MainWindow(QWidget):
    def __init__(self):
        super().__init__()
        # === УСТАНОВКА ИКОНКИ ОКНА ===
        icon_path = get_resource_path("logo.png")
        if os.path.exists(icon_path):
            self.setWindowIcon(QIcon(icon_path))
        else:
            print("Предупреждение: logo.png не найден")
        self.setWindowTitle("RemzonaDownloader")
        self.resize(950, 720)

        self.reader = UniversalTableReader()
        self.worker = None
        self.compress_worker = None
        self.splitter_worker = None
        self.converter_worker = None
        self.compressor = Compressor()
        self.compress_selected_folder = None

        main_layout = QVBoxLayout()
        self.setLayout(main_layout)

        self.tabs = QTabWidget()
        main_layout.addWidget(self.tabs)

        self.tab_download = self._create_download_tab()
        self.tab_compress = self._create_compress_tab()
        self.tab_split = self._create_split_tab()   
        self.tab_converter = self._create_converter_tab()

        self.tabs.addTab(self.tab_download, "Скачать файлы по ссылкам")
        self.tabs.addTab(self.tab_compress, "Упаковать файлы в архив")
        self.tabs.addTab(self.tab_split, "Разбивка файла")
        self.tabs.addTab(self.tab_converter, "Конвертер Excel")

        self.tabs.setCurrentIndex(0)

        # ====================== ОБЩАЯ ФУНКЦИЯ ДЛЯ ЗАГОЛОВКА ======================
    def create_header(self, title_text):
        header_layout = QHBoxLayout()
        logo = QLabel()
        logo.setAlignment(Qt.AlignmentFlag.AlignCenter)

        pixmap = QPixmap(get_resource_path("logo.png"))
        if pixmap.isNull():
            logo.setText("🖼️")
            logo.setFont(QFont("Segoe UI", 28))
        else:
            scaled = pixmap.scaled(65, 65, Qt.AspectRatioMode.KeepAspectRatio, 
                                 Qt.TransformationMode.SmoothTransformation)
            logo.setPixmap(scaled)

        title = QLabel(title_text)
        title.setFont(QFont("Segoe UI", 14, QFont.Weight.Bold))
        title.setAlignment(Qt.AlignmentFlag.AlignVCenter)

        header_layout.addStretch()
        header_layout.addWidget(logo)
        header_layout.addSpacing(15)
        header_layout.addWidget(title)
        header_layout.addStretch()
        return header_layout
    # ============================================================================



        # ====================== CONVERTER TAB ======================
    def _create_converter_tab(self):
        tab = QWidget()
        layout = QVBoxLayout()
        layout.addLayout(self.create_header("Конвертер Excel → TXT / CSV"))
        layout.addSpacing(20)

        btn_layout = QHBoxLayout()
        self.conv_folder_btn = QPushButton("Выбрать папку с Excel файлами")
        self.conv_folder_btn.setFixedWidth(280)
        self.conv_folder_btn.clicked.connect(self.conv_select_folder)
        btn_layout.addStretch()
        btn_layout.addWidget(self.conv_folder_btn)
        btn_layout.addStretch()
        layout.addLayout(btn_layout)
        layout.addSpacing(20)

        format_layout = QHBoxLayout()
        format_label = QLabel("Формат вывода:")
        self.conv_radio_txt = QRadioButton("Текст Юникод (.txt)")
        self.conv_radio_csv = QRadioButton("CSV с разделителем ; (.csv)")
        self.conv_radio_txt.setChecked(True)

        format_layout.addStretch()
        format_layout.addWidget(format_label)
        format_layout.addWidget(self.conv_radio_txt)
        format_layout.addWidget(self.conv_radio_csv)
        format_layout.addStretch()
        layout.addLayout(format_layout)
        layout.addSpacing(30)

        start_layout = QHBoxLayout()
        self.conv_start_btn = QPushButton("Начать конвертацию")
        self.conv_start_btn.setFixedWidth(260)
        self.conv_start_btn.clicked.connect(self.conv_on_start_clicked)
        start_layout.addStretch()
        start_layout.addWidget(self.conv_start_btn)
        start_layout.addStretch()
        layout.addLayout(start_layout)
        layout.addSpacing(20)

        self.conv_progress = QProgressBar()
        layout.addWidget(self.conv_progress)
        layout.addSpacing(10)

        self.conv_log = QTextEdit()
        self.conv_log.setReadOnly(True)
        layout.addWidget(self.conv_log)

        tab.setLayout(layout)
        return tab

    def conv_select_folder(self):
        directory = QFileDialog.getExistingDirectory(self, "Выберите папку с Excel файлами")
        if directory:
            self.conv_folder = directory
            self.conv_log.append(f"Выбрана папка: {directory}")

    def conv_on_start_clicked(self):
        if self.converter_worker and self.converter_worker.isRunning():
            self.converter_worker.cancel()
            self.conv_start_btn.setText("Начать конвертацию")
            return

        if not hasattr(self, 'conv_folder'):
            QMessageBox.warning(self, "Ошибка", "Сначала выберите папку с Excel файлами!")
            return

        choice = "1" if self.conv_radio_txt.isChecked() else "2"

        self.conv_log.clear()
        self.conv_log.append("Запуск конвертации...\n")
        self.conv_progress.setValue(0)
        self.conv_progress.setMaximum(100)

        self.converter_worker = ConverterWorker(self.conv_folder, choice)
        self.converter_worker.log.connect(self.conv_log.append)
        self.converter_worker.progress.connect(self.conv_progress.setValue)
        self.converter_worker.result.connect(lambda c: self.conv_log.append(f"\n✅ Обработано файлов: {c}"))
        self.converter_worker.finished.connect(self.conv_on_finished)

        self.converter_worker.start()
        self.conv_start_btn.setText("Отмена")

    def conv_on_finished(self):
        self.conv_start_btn.setText("Начать конвертацию")

   # ====================== DOWNLOAD TAB ======================
    def _create_download_tab(self):
        tab = QWidget()
        layout = QVBoxLayout()
        layout.addLayout(self.create_header("Скачивание изображений по ссылкам"))
        layout.addSpacing(20)

        self.download_stack_layout = QVBoxLayout()
        layout.addLayout(self.download_stack_layout)

        self.download_start_screen = self._create_download_start_screen()
        self.download_preview_screen = self._create_download_preview_screen()
        self.download_process_screen = self._create_download_process_screen()

        self.download_screens = [self.download_start_screen, self.download_preview_screen, self.download_process_screen]
        self.download_stack_layout.addWidget(self.download_start_screen)

        tab.setLayout(layout)
        return tab


        # ============================================================================

        # Стек экранов (старт / предпросмотр / процесс)
        self.download_stack_layout = QVBoxLayout()
        layout.addLayout(self.download_stack_layout)

        self.download_start_screen = self._create_download_start_screen()
        self.download_preview_screen = self._create_download_preview_screen()
        self.download_process_screen = self._create_download_process_screen()

        self.download_screens = [
            self.download_start_screen,
            self.download_preview_screen,
            self.download_process_screen
        ]

        self.download_stack_layout.addWidget(self.download_start_screen)

        tab.setLayout(layout)
        return tab

    # (Все методы download_... оставлены как у тебя)
    def _create_download_start_screen(self):
        screen = QWidget()
        layout = QVBoxLayout()
        btn = QPushButton("Выбрать файл")
        btn.setFixedWidth(150)
        btn.clicked.connect(self.download_open_file)
        layout.addWidget(btn, alignment=Qt.AlignmentFlag.AlignCenter)

        placeholder = QLabel("Поддерживаемые файлы в формате .csv, .xls, .xlsx")
        placeholder.setStyleSheet("color: gray; font-size: 14px;")
        placeholder_layout = QHBoxLayout()
        placeholder_layout.addStretch()
        placeholder_layout.addWidget(placeholder)
        placeholder_layout.addStretch()
        layout.addLayout(placeholder_layout)
        layout.addStretch()
        screen.setLayout(layout)
        return screen

    def _create_download_preview_screen(self):
        screen = QWidget()
        layout = QVBoxLayout()
        top = QHBoxLayout()
        self.download_back_btn = QPushButton("Назад")
        self.download_back_btn.clicked.connect(self.download_go_back)
        self.download_forward_btn = QPushButton("Вперёд")
        self.download_forward_btn.setEnabled(False)
        self.download_forward_btn.clicked.connect(self.download_process_data)
        top.addWidget(self.download_back_btn)
        top.addStretch()
        top.addWidget(self.download_forward_btn)
        layout.addSpacing(10)
        layout.addLayout(top)
        layout.addSpacing(20)
        self.download_table = QTableView()
        layout.addWidget(self.download_table)
        screen.setLayout(layout)
        return screen

    def _create_download_process_screen(self):
        screen = QWidget()
        layout = QVBoxLayout()
        top_layout = QHBoxLayout()
        self.download_back_process_btn = QPushButton("Назад")
        self.download_back_process_btn.setFixedWidth(100)
        self.download_back_process_btn.clicked.connect(self.download_go_back_to_preview)
        top_layout.addWidget(self.download_back_process_btn)
        top_layout.addStretch()
        layout.addLayout(top_layout)
        layout.addSpacing(20)

        center_layout = QHBoxLayout()
        self.download_dir_button = QPushButton("Выбрать папку для скачивания")
        self.download_dir_button.setFixedWidth(260)
        self.download_dir_button.clicked.connect(self.download_select_directory)
        center_layout.addStretch()
        center_layout.addWidget(self.download_dir_button)
        center_layout.addStretch()
        layout.addLayout(center_layout)
        layout.addSpacing(20)

        delim_layout = QHBoxLayout()
        delim_label = QLabel("Разделитель ссылок:")
        self.download_delimiter_input = QLineEdit()
        self.download_delimiter_input.setFixedWidth(40)
        delim_layout.addStretch()
        delim_layout.addWidget(delim_label)
        delim_layout.addSpacing(10)
        delim_layout.addWidget(self.download_delimiter_input)
        delim_layout.addStretch()
        layout.addLayout(delim_layout)
        layout.addSpacing(20)

        headers_layout = QHBoxLayout()
        headers_label = QLabel("Пропустить первую строку (заголовки)?")
        self.download_radio_skip = QRadioButton("Пропустить")
        self.download_radio_no_skip = QRadioButton("Не пропускать")
        self.download_headers_group = QButtonGroup()
        self.download_headers_group.addButton(self.download_radio_skip)
        self.download_headers_group.addButton(self.download_radio_no_skip)
        self.download_radio_skip.setChecked(True)
        self.download_radio_skip.toggled.connect(self.download_update_headers)
        headers_layout.addStretch()
        headers_layout.addWidget(headers_label)
        headers_layout.addSpacing(10)
        headers_layout.addWidget(self.download_radio_skip)
        headers_layout.addWidget(self.download_radio_no_skip)
        headers_layout.addStretch()
        layout.addLayout(headers_layout)
        layout.addSpacing(20)

        start_layout = QHBoxLayout()
        self.download_start_button = QPushButton("Начать")
        self.download_start_button.setFixedWidth(260)
        self.download_start_button.clicked.connect(self.download_on_start_cancel_clicked)
        start_layout.addStretch()
        start_layout.addWidget(self.download_start_button)
        start_layout.addStretch()
        layout.addLayout(start_layout)
        layout.addSpacing(20)

        self.download_progress = QProgressBar()
        self.download_progress.setValue(0)
        self.download_progress.setEnabled(False)
        layout.addWidget(self.download_progress)
        layout.addSpacing(10)

        self.download_log_output = QTextEdit()
        self.download_log_output.setReadOnly(True)
        layout.addWidget(self.download_log_output)

        screen.setLayout(layout)
        return screen

    # -------- DOWNLOAD TAB METHODS --------
    def download_open_file(self):
        """Open file dialog and load file."""
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Файл",
            "",
            "All (*);;CSV (*.csv);;Excel (*.xlsx)"
        )

        if not path:
            return

        try:
            self.reader.read_file(path)
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", str(e))
            return

        self.download_show_preview()

    def download_show_preview(self):
        """Show preview of loaded file."""
        preview = self.reader.data_frame.head(10)

        self.download_model = PandasModel(preview)
        self.download_table.setModel(self.download_model)

        header = self.download_table.horizontalHeader()
        header.setContextMenuPolicy(Qt.CustomContextMenu)
        header.customContextMenuRequested.connect(self.download_open_menu)

        self._switch_download_screen(self.download_preview_screen)

    def download_open_menu(self, pos):
        """Open context menu for column selection."""
        header = self.download_table.horizontalHeader()
        col = header.logicalIndexAt(pos)

        if col < 0:
            return

        menu = QMenu(self)

        url_action = menu.addAction("Использовать как ссылки")
        name_action = menu.addAction("Использовать как имена файлов")
        brand_action = menu.addAction("Использовать как бренд")

        action = menu.exec(header.mapToGlobal(pos))

        if action is url_action:
            self.download_model.url_col = col
            self.reader.urls_column_index = col
            self.download_forward_btn.setEnabled(True)

        elif action is name_action:
            self.download_model.name_col = col
            self.reader.filenames_column_index = col

        elif action is brand_action:
            self.download_model.brand_col = col
            self.reader.brands_column_index = col

        self.download_table.viewport().update()

    def download_process_data(self):
        """Move to process screen."""
        self.reader.urls_column_index = self.download_model.url_col
        self.reader.filenames_column_index = self.download_model.name_col
        self.reader.brands_column_index = self.download_model.brand_col
        self._switch_download_screen(self.download_process_screen)

    def download_go_back(self):
        """Go back to start screen."""
        self.reader = UniversalTableReader()
        self.download_table.setModel(None)
        self.download_forward_btn.setEnabled(False)

        self._switch_download_screen(self.download_start_screen)

    def download_go_back_to_preview(self):
        """Go back to preview screen."""
        self._switch_download_screen(self.download_preview_screen)

    def download_select_directory(self):
        """Select download directory."""
        directory = QFileDialog.getExistingDirectory(
            self,
            "Выберите папку"
        )

        if not directory:
            return

        self.reader.download_dir = directory
        self.download_log_output.append(
            f"Папка для скачивания: {directory}"
        )

    def download_on_start_cancel_clicked(self):
        """Handle start/cancel button click."""
        if self.worker is not None and self.worker.is_running:
            self.download_cancel_processing()
        else:
            self.download_start_processing()

    def download_start_processing(self):
        """Start file download process."""
        try:
            delimiter = self.download_delimiter_input.text().strip()

            if delimiter:
                self.reader.urls_delimiter = delimiter

            if not self.reader.download_dir:
                QMessageBox.warning(
                    self,
                    "Ошибка",
                    "Сначала выберите папку"
                )
                return

            data = self.reader.preparing_data()

            data = [
                {
                    "url": item["url"],
                    "file_name": item["file_name"],
                    "brand": item.get("brand")
                }
                for item in data
            ]

            self.download_progress.setEnabled(True)
            self.download_progress.setMaximum(len(data))
            self.download_progress.setValue(0)

            self.download_log_output.append(
                "Начинаю скачивание файлов по ссылкам...\n"
            )

            # Create worker
            self.worker = DownloadWorker(data, self.reader.download_dir)

            # Connect signals
            self.worker.progress.connect(self.download_update_progress)
            self.worker.log_item.connect(self.download_add_log_item)
            self.worker.result.connect(self.download_show_result)
            self.worker.cancelled.connect(self.download_on_cancelled)
            self.worker.log.connect(self.download_add_error_log)
            self.worker.finished.connect(self.download_on_finished)

            # Start worker
            self.worker.start()

            # Update button
            self.download_start_button.setText("Отмена")

            # Disable controls
            self.download_dir_button.setEnabled(False)
            self.download_delimiter_input.setEnabled(False)
            self.download_radio_skip.setEnabled(False)
            self.download_radio_no_skip.setEnabled(False)

        except Exception as e:
            QMessageBox.critical(self, "Ошибка", str(e))

    def download_cancel_processing(self):
        """Cancel ongoing download."""
        if self.worker is not None and self.worker.is_running:
            self.download_log_output.append(
                "\n<span style='color:orange'>Отмена скачивания...</span>"
            )
            self.worker.cancel_download()

    def download_on_finished(self):
        """Called when worker finishes."""
        self.download_start_button.setText("Начать")

        self.download_dir_button.setEnabled(True)
        self.download_delimiter_input.setEnabled(True)
        self.download_radio_skip.setEnabled(True)
        self.download_radio_no_skip.setEnabled(True)

    def download_on_cancelled(self):
        """Called when download is cancelled."""
        self.download_log_output.append(
            "\n<span style='color:orange'>Загрузка отменена пользователем."
            "</span>"
        )

    def download_add_log_item(self, url, success):
        """Add log item."""
        color = "green" if success else "red"
        self.download_log_output.append(
            f'<span style="color:{color}">{url}</span>'
        )

    def download_update_progress(self, value):
        """Update progress bar."""
        self.download_progress.setValue(value)

    def download_show_result(self, total, errors_count):
        """Show download result."""
        success = total - errors_count

        self.download_log_output.append("")

        self.download_log_output.append(
            f'<span style="color:black">'
            f'Готово. Успешно скачано файлов: {success} из {total}. '
            f'Ошибок: {errors_count}.'
            f'</span>'
        )

    def download_add_error_log(self, text):
        """Add error log."""
        self.download_log_output.append(
            f'<span style="color:black">{text}</span>'
        )

    def download_update_headers(self):
        """Update headers option."""
        self.reader.headers_column = self.download_radio_skip.isChecked()

    def _switch_download_screen(self, new_screen):
        """Switch download tab screen."""
        # Clear existing layout
        while self.download_stack_layout.count():
            item = self.download_stack_layout.takeAt(0)
            if item.widget():
                item.widget().hide()

        # Add new screen
        self.download_stack_layout.addWidget(new_screen)
        new_screen.show()

    # ========================================================================
    # COMPRESS TAB (оставлен как был)
    # ========================================================================
    def _create_compress_tab(self):
        tab = QWidget()
        layout = QVBoxLayout()
        layout.addLayout(self.create_header("Упаковка файлов в архив"))
        layout.addSpacing(20)

        self.compress_stack_layout = QVBoxLayout()
        layout.addLayout(self.compress_stack_layout)

        self.compress_start_screen = self._create_compress_start_screen()
        self.compress_params_screen = self._create_compress_params_screen()
        self.compress_stack_layout.addWidget(self.compress_start_screen)

        tab.setLayout(layout)
        return tab

    def _create_compress_start_screen(self):
        """Create start screen for compress tab."""
        screen = QWidget()
        layout = QVBoxLayout()

        btn = QPushButton("Выбрать папку")
        btn.setFixedWidth(150)
        btn.clicked.connect(self.compress_select_folder)

        layout.addWidget(btn, alignment=Qt.AlignmentFlag.AlignCenter)

        placeholder = QLabel("Выберите папку с файлами для архивирования")
        placeholder.setStyleSheet("color: gray; font-size: 14px;")
        layout.addWidget(placeholder, alignment=Qt.AlignmentFlag.AlignCenter)

        layout.addStretch()

        screen.setLayout(layout)
        return screen

    def _create_compress_params_screen(self):
        """Create parameters screen for compress tab with progress and logs."""
        screen = QWidget()
        layout = QVBoxLayout()

        # -------- BACK BUTTON (TOP LEFT) --------
        top_layout = QHBoxLayout()

        self.compress_back_params_btn = QPushButton("Назад")
        self.compress_back_params_btn.setFixedWidth(100)
        self.compress_back_params_btn.clicked.connect(self.compress_go_back)

        top_layout.addWidget(self.compress_back_params_btn)
        top_layout.addStretch()

        layout.addLayout(top_layout)
        layout.addSpacing(20)

        # -------- ARCHIVE NAME INPUT (CENTER) --------
        archive_name_layout = QHBoxLayout()

        archive_name_label = QLabel("Имя архива(ов):")

        self.compress_archive_name_input = QLineEdit()
        self.compress_archive_name_input.setFixedWidth(200)
        self.compress_archive_name_input.setPlaceholderText("archive")

        archive_name_layout.addStretch()
        archive_name_layout.addWidget(archive_name_label)
        archive_name_layout.addSpacing(10)
        archive_name_layout.addWidget(self.compress_archive_name_input)
        archive_name_layout.addStretch()

        layout.addLayout(archive_name_layout)
        layout.addSpacing(20)

        # -------- MAX FILES PER ARCHIVE INPUT (CENTER) --------
        max_files_layout = QHBoxLayout()

        max_files_label = QLabel("Макс. количество файлов в архиве:")

        self.compress_max_files_input = QSpinBox()
        self.compress_max_files_input.setFixedWidth(100)
        self.compress_max_files_input.setMinimum(1)
        self.compress_max_files_input.setMaximum(10000)
        self.compress_max_files_input.setValue(100)

        max_files_layout.addStretch()
        max_files_layout.addWidget(max_files_label)
        max_files_layout.addSpacing(10)
        max_files_layout.addWidget(self.compress_max_files_input)
        max_files_layout.addStretch()

        layout.addLayout(max_files_layout)
        layout.addSpacing(30)

        # -------- START BUTTON --------
        start_layout = QHBoxLayout()

        self.compress_start_button = QPushButton("Начать")
        self.compress_start_button.setFixedWidth(260)
        self.compress_start_button.clicked.connect(
            self.compress_on_start_cancel_clicked
        )

        start_layout.addStretch()
        start_layout.addWidget(self.compress_start_button)
        start_layout.addStretch()

        layout.addLayout(start_layout)
        layout.addSpacing(20)

        # -------- PROGRESS --------
        self.compress_progress = QProgressBar()
        self.compress_progress.setValue(0)
        self.compress_progress.setEnabled(False)

        layout.addWidget(self.compress_progress)
        layout.addSpacing(10)

        # -------- LOG --------
        self.compress_log_output = QTextEdit()
        self.compress_log_output.setReadOnly(True)

        layout.addWidget(self.compress_log_output)

        screen.setLayout(layout)
        return screen

    # -------- COMPRESS TAB METHODS --------
    def compress_select_folder(self):
        """Open folder selection dialog and set up compression."""
        directory = QFileDialog.getExistingDirectory(
            self,
            "Выберите папку с файлами"
        )

        if not directory:
            return

        self.compress_selected_folder = directory
        self.compressor.path2dir = directory

        self._switch_compress_screen(self.compress_params_screen)

    def compress_go_back(self):
        """Go back to start screen."""
        self.compress_selected_folder = None
        self.compressor.path2dir = None
        self.compress_archive_name_input.clear()
        self.compress_max_files_input.setValue(10000)
        self.compress_log_output.clear()
        self.compress_progress.setValue(0)

        self._switch_compress_screen(self.compress_start_screen)

    def compress_on_start_cancel_clicked(self):
        """Handle start/cancel button click."""

        if self.compress_worker is not None and self.compress_worker.is_running:
            self.compress_cancel_processing()
        else:
            self.compress_start_processing()

    def compress_on_cancelled(self):
        self.compress_log_output.append(
            "\n<span style='color:orange'>Архивирование отменено пользователем.</span>"
        )

    def compress_start_processing(self):
        """Start file compression process."""
        try:
            archive_name = self.compress_archive_name_input.text().strip()
            max_files = self.compress_max_files_input.value()
            if not archive_name:
                QMessageBox.warning(
                    self,
                    "Ошибка",
                    "Пожалуйста, введите имя архива"
                )
                return

            if max_files <= 0:
                QMessageBox.warning(
                    self,
                    "Ошибка",
                    "Количество файлов должно быть больше нуля"
                )
                return

            if not self.compress_selected_folder:
                QMessageBox.warning(
                    self,
                    "Ошибка",
                    "Папка не выбрана"
                )
                return

            # Store parameters for worker
            archive_name_value = archive_name
            max_files_value = max_files

            self.compress_progress.setEnabled(True)
            self.compress_progress.setValue(0)

            self.compress_log_output.clear()
            self.compress_log_output.append(
                f"Начинаю архивирование папки: {self.compress_selected_folder}\n"
            )
            self.compress_log_output.append(
                f"Имя архива: {archive_name_value}\n"
            )
            self.compress_log_output.append(
                f"Макс. файлов в архиве: {max_files_value}\n\n"
            )

            # Create new compressor instance for this operation
            compressor = Compressor()
            compressor.path2dir = self.compress_selected_folder
            compressor.archive_name = archive_name_value
            compressor.pack = int(max_files)

            # Create worker
            self.compress_worker = CompressorWorker(compressor)

            # Connect signals
            self.compress_worker.progress.connect(self.compress_update_progress)
            self.compress_worker.log.connect(self.compress_add_log)
            self.compress_worker.result.connect(self.compress_show_result)
            self.compress_worker.finished.connect(self.compress_on_finished)
            self.compress_worker.cancelled.connect(self.compress_on_cancelled)

            # Start worker
            self.compress_worker.start()

            # Update button
            self.compress_start_button.setText("Отмена")

            # Disable controls
            self.compress_archive_name_input.setEnabled(False)
            self.compress_max_files_input.setEnabled(False)
            self.compress_back_params_btn.setEnabled(False)

        except Exception as e:
            QMessageBox.critical(self, "Ошибка", str(e))

    def compress_cancel_processing(self):
        """Cancel ongoing compression."""
        if self.compress_worker and self.compress_worker.is_running:
            self.compress_log_output.append(
                "\n<span style='color:orange'>Отмена архивирования...</span>"
            )
            self.compress_worker.cancel_compression()

    def compress_on_finished(self):
        """Called when worker finishes."""
        self.compress_start_button.setText("Начать")

        self.compress_archive_name_input.setEnabled(True)
        self.compress_max_files_input.setEnabled(True)
        self.compress_back_params_btn.setEnabled(True)

    def compress_update_progress(self, value):
        """Update progress bar."""
        if self.compress_worker and self.compress_worker.compressor.all_files_count:
            max_value = self.compress_worker.compressor.all_files_count
            self.compress_progress.setMaximum(max_value)
            self.compress_progress.setValue(value)

    def compress_add_log(self, text):
        """Add log message."""
        self.compress_log_output.append(
            f'<span style="color:black">{text}</span>'
        )

    def compress_show_result(self, total, errors_count):
        """Show compression result."""
        self.compress_log_output.append("")

        if errors_count == 0:
            self.compress_log_output.append(
                f'<span style="color:green">'
                f'Готово. Архивировано файлов: {total}.'
                f'</span>'
            )
        else:
            self.compress_log_output.append(
                f'<span style="color:red">'
                f'Ошибка при архивировании. Обработано файлов: {total}. '
                f'Ошибок: {errors_count}.'
                f'</span>'
            )

    def _switch_compress_screen(self, new_screen):
        """Switch compress tab screen."""
        # Clear existing layout
        while self.compress_stack_layout.count():
            item = self.compress_stack_layout.takeAt(0)
            if item.widget():
                item.widget().hide()

        # Add new screen
        self.compress_stack_layout.addWidget(new_screen)
        new_screen.show()


# ====================== SPLIT TAB ======================
    def _create_split_tab(self):
        tab = QWidget()
        layout = QVBoxLayout()
        layout.addLayout(self.create_header("Разбивка большого Excel-файла на части"))
        layout.addSpacing(25)

        file_layout = QHBoxLayout()
        self.split_file_btn = QPushButton("Выбрать Excel файл")
        self.split_file_btn.setFixedWidth(240)
        self.split_file_btn.clicked.connect(self.split_select_file)
        file_layout.addStretch()
        file_layout.addWidget(self.split_file_btn)
        file_layout.addStretch()
        layout.addLayout(file_layout)
        layout.addSpacing(10)

        out_layout = QHBoxLayout()
        self.split_output_btn = QPushButton("Выбрать папку для сохранения")
        self.split_output_btn.setFixedWidth(240)
        self.split_output_btn.clicked.connect(self.split_select_output_dir)
        out_layout.addStretch()
        out_layout.addWidget(self.split_output_btn)
        out_layout.addStretch()
        layout.addLayout(out_layout)
        layout.addSpacing(15)

        chunk_layout = QHBoxLayout()
        chunk_label = QLabel("Строк в одном файле:")
        self.split_chunk_spin = QSpinBox()
        self.split_chunk_spin.setFixedWidth(140)
        self.split_chunk_spin.setMinimum(100)
        self.split_chunk_spin.setMaximum(1000000)
        self.split_chunk_spin.setValue(150000)
        self.split_chunk_spin.setSingleStep(10000)
        chunk_layout.addStretch()
        chunk_layout.addWidget(chunk_label)
        chunk_layout.addWidget(self.split_chunk_spin)
        chunk_layout.addStretch()
        layout.addLayout(chunk_layout)
        layout.addSpacing(30)

        start_layout = QHBoxLayout()
        self.split_start_btn = QPushButton("Начать разбивку")
        self.split_start_btn.setFixedWidth(260)
        self.split_start_btn.clicked.connect(self.split_on_start_clicked)
        start_layout.addStretch()
        start_layout.addWidget(self.split_start_btn)
        start_layout.addStretch()
        layout.addLayout(start_layout)
        layout.addSpacing(20)

        self.split_progress = QProgressBar()
        self.split_progress.setValue(0)
        layout.addWidget(self.split_progress)
        layout.addSpacing(10)

        self.split_log = QTextEdit()
        self.split_log.setReadOnly(True)
        layout.addWidget(self.split_log)

        tab.setLayout(layout)
        return tab



    def split_select_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Выберите Excel файл", "", "Excel (*.xlsx *.xls)")
        if path:
            self.split_input_file = path
            self.split_log.append(f"Выбран файл: {path}")

    def split_select_output_dir(self):
        directory = QFileDialog.getExistingDirectory(self, "Выберите папку для сохранения")
        if directory:
            self.split_output_dir = directory
            self.split_log.append(f"Папка сохранения: {directory}")

    def split_on_start_clicked(self):
        if self.splitter_worker and self.splitter_worker.isRunning():
            self.splitter_worker.cancel()
            self.split_start_btn.setText("Начать разбивку")
            return

        if not hasattr(self, 'split_input_file') or not hasattr(self, 'split_output_dir'):
            QMessageBox.warning(self, "Ошибка", "Сначала выберите файл и папку!")
            return

        chunk_size = self.split_chunk_spin.value()

        self.split_log.clear()
        self.split_log.append("Запуск разбивки...\n")
        self.split_progress.setValue(0)
        self.split_progress.setMaximum(100)

        self.splitter_worker = SplitterWorker(
            self.split_input_file,
            self.split_output_dir,
            chunk_size
        )

        self.splitter_worker.log.connect(self.split_log.append)
        self.splitter_worker.progress.connect(self.split_progress.setValue)
        self.splitter_worker.result.connect(lambda count: self.split_log.append(f"\n✅ Создано файлов: {count}"))
        self.splitter_worker.finished.connect(self.split_on_finished)

        self.splitter_worker.start()
        self.split_start_btn.setText("Отмена")

    def split_on_finished(self):
        self.split_progress.setValue(100)          # ← ИСПРАВЛЕНИЕ
        self.split_start_btn.setText("Начать разбивку")